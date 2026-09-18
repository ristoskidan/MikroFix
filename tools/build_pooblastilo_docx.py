from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "pravni" / "POOBLASTILO_ZA_PREGLED_IN_PREVZEM_NAPRAV.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_run(run, size=10.5, bold=False, italic=False, color="000000"):
    run.font.name = "Liberation Sans"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Liberation Sans")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_label_value(paragraph, label, value, size=10.5):
    configure_run(paragraph.add_run(label), size=size, bold=True)
    configure_run(paragraph.add_run(value), size=size)


def add_numbered_item(doc, number, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.23)
    paragraph.paragraph_format.first_line_indent = Inches(-0.23)
    paragraph.paragraph_format.space_after = Pt(2.3)
    paragraph.paragraph_format.line_spacing = 1.0
    configure_run(paragraph.add_run(f"{number}. "), size=9.7, bold=True)
    configure_run(paragraph.add_run(text), size=9.7)
    return paragraph


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.52)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.03

    title_style = doc.styles["Title"]
    title_style.font.name = "Liberation Sans"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    title_style.font.size = Pt(17)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    remove_paragraph_border(title_style)

    heading_style = doc.styles["Heading 1"]
    heading_style.font.name = "Liberation Sans"
    heading_style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    heading_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    heading_style.font.size = Pt(11.5)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(0, 0, 0)
    heading_style.paragraph_format.space_before = Pt(4)
    heading_style.paragraph_format.space_after = Pt(2)

    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(4.9)
    header_table.columns[1].width = Inches(2.3)
    left, right = header_table.rows[0].cells
    left.width = Inches(4.9)
    right.width = Inches(2.3)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(left, top=20, start=0, bottom=50, end=80)
    set_cell_margins(right, top=20, start=80, bottom=50, end=0)
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    configure_run(p.add_run("MikroFix"), size=15, bold=True, color="2F6FAE")
    configure_run(p.add_run("\nRistoski Daniel s.p."), size=9.3, bold=True)
    configure_run(p.add_run("\nParadiž 73, 2282 Cirkulane"), size=8.8)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    configure_run(p.add_run("Matična št.: 7565178000\n"), size=8.8)
    configure_run(p.add_run("Davčna št.: 84514841\n"), size=8.8)
    configure_run(p.add_run("Dokument: MF-POO-____/____"), size=8.8, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(7)
    title.paragraph_format.space_after = Pt(2)
    title.add_run("POOBLASTILO ZA PREGLED IN PREVZEM NAPRAV")
    remove_paragraph_border(title)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    configure_run(subtitle.add_run("Omejeno pooblastilo za postopke odkupa premičnih elektronskih naprav"), size=9.5, italic=True)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(4)
    add_label_value(intro, "Pooblastitelj: ", "MikroFix, popravilo elektronskih in telekomunikacijskih naprav, Ristoski Daniel s.p., Paradiž 73, 2282 Cirkulane, ki ga zastopa podjetnik Daniel Ristoski.", size=9.7)

    doc.add_paragraph("Podatki pooblaščene osebe", style="Heading 1")
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.65)
    table.columns[1].width = Inches(5.55)
    set_table_borders(table)
    fields = [
        ("Ime in priimek", "________________________________________________________________"),
        ("Naslov prebivališča", "________________________________________________________________"),
        ("Datum rojstva", "________________________________________________________________"),
        ("Kontakt", "________________________________________________________________"),
    ]
    for idx, (label, line) in enumerate(fields):
        label_cell, value_cell = table.rows[idx].cells
        label_cell.width = Inches(1.65)
        value_cell.width = Inches(5.55)
        label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        value_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(label_cell, top=55, start=95, bottom=55, end=95)
        set_cell_margins(value_cell, top=55, start=95, bottom=55, end=95)
        set_cell_shading(label_cell, "E9EFF5")
        label_cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        value_cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        configure_run(label_cell.paragraphs[0].add_run(label), size=9.2, bold=True)
        configure_run(value_cell.paragraphs[0].add_run(line), size=9.2)

    identity = doc.add_paragraph()
    identity.paragraph_format.space_before = Pt(2)
    identity.paragraph_format.space_after = Pt(2)
    configure_run(identity.add_run("Pooblaščena oseba se pri opravljanju pooblastila na zahtevo izkaže z veljavnim osebnim dokumentom. Kopija osebnega dokumenta ni priloga tega pooblastila."), size=8.8, italic=True)

    doc.add_paragraph("Obseg pooblastila", style="Heading 1")
    scope_items = [
        "na območju Republike Slovenije pregleda napravo, ki je uvrščena na veljavni odkupni seznam MikroFixa, ter preveri njeno dejansko stanje in priloženo opremo;",
        "z aplikacijama MikroFIX Service in MicroDetect zajame tehnične ter identifikacijske podatke naprave in pripravi zapis pregleda;",
        "zbere potrebne podatke prodajalca, vizualno preveri njegovo identiteto in pridobi izjavo prodajalca o lastništvu naprave;",
        "pripravi osnutek odkupne pogodbe, ga v celoti pokaže prodajalcu in pridobi prodajalčev podpis;",
        "po izrecni odobritvi in podpisu Daniela Ristoski ter po potrditvi plačila prevzame napravo, pripadajočo opremo in izročene dokumente;",
        "napravo varno prevaža ter jo brez nepotrebnega odlašanja preda v delavnico oziroma pisarno MikroFixa ob evidentiranem prenosu skrbništva.",
    ]
    for idx, text in enumerate(scope_items, 1):
        add_numbered_item(doc, idx, text)

    doc.add_paragraph("Omejitve pooblastila", style="Heading 1")
    restrictions = doc.add_paragraph()
    restrictions.paragraph_format.space_after = Pt(3)
    restrictions.paragraph_format.line_spacing = 1.0
    configure_run(restrictions.add_run("Pooblaščena oseba brez ločenega pisnega pooblastila ne sme "), size=9.5, bold=True)
    configure_run(restrictions.add_run("v imenu MikroFixa podpisati odkupne pogodbe ali drugače dokončno zavezati pooblastitelja, spremeniti odkupne cene ali pravila, opraviti ali obljubiti plačila, dostopati do bančnega računa ali digitalnih potrdil, prevzeti naprave pred potrditvijo plačila, prenesti pooblastila na tretjo osebo, naprave uporabljati, prodati, zastaviti, popravljati ali razstaviti, razen kolikor je nujno za dovoljen pregled."), size=9.5)

    conditions = doc.add_paragraph()
    conditions.paragraph_format.space_after = Pt(3)
    conditions.paragraph_format.line_spacing = 1.0
    configure_run(conditions.add_run("Veljavnost in preklic: "), size=9.5, bold=True)
    configure_run(conditions.add_run("Pooblastilo velja od __________________ do __________________, vendar največ eno leto. Pooblastitelj ga lahko kadarkoli pisno prekliče. Po prenehanju veljavnosti mora pooblaščena oseba prenehati z vsemi postopki in vrniti vse prejete predmete ter dokumente."), size=9.5)

    clarification = doc.add_paragraph()
    clarification.paragraph_format.space_after = Pt(5)
    clarification.paragraph_format.line_spacing = 1.0
    configure_run(clarification.add_run("To pooblastilo samo po sebi ne predstavlja pogodbe o zaposlitvi, pravice do plačila ali dovoljenja za izdajanje računov. Podlaga za občasno brezplačno pomoč in potrebne evidence se uredijo ločeno."), size=8.7, italic=True)

    verification = doc.add_paragraph()
    verification.paragraph_format.space_after = Pt(6)
    add_label_value(verification, "Kontakt za preveritev veljavnosti pooblastila: ", "____________________________________________", size=9.3)

    signature_table = doc.add_table(rows=4, cols=2)
    signature_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    signature_table.autofit = False
    signature_table.columns[0].width = Inches(3.45)
    signature_table.columns[1].width = Inches(3.45)
    left, right = signature_table.rows[0].cells
    for cell in (left, right):
        set_cell_margins(cell, top=20, start=0, bottom=20, end=0)
    add_label_value(left.paragraphs[0], "Kraj: ", "____________________________", size=9.3)
    add_label_value(right.paragraphs[0], "Datum: ", "__________________________", size=9.3)
    signature_table.cell(1, 0).merge(signature_table.cell(1, 1))
    blank = signature_table.cell(1, 0).paragraphs[0]
    blank.paragraph_format.space_after = Pt(18)

    headers = signature_table.rows[2].cells
    configure_run(headers[0].paragraphs[0].add_run("POOBLASTITELJ"), size=9.2, bold=True)
    configure_run(headers[1].paragraphs[0].add_run("POOBLAŠČENA OSEBA"), size=9.2, bold=True)
    headers[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    headers[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    signatures = signature_table.rows[3].cells
    for cell in signatures:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].paragraph_format.space_before = Pt(23)
        configure_run(cell.paragraphs[0].add_run("__________________________________"), size=9.3)
        configure_run(cell.add_paragraph().add_run("podpis"), size=8.3, italic=True)
        cell.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[-1].paragraph_format.space_after = Pt(0)

    # Keep the form visually clean without visible signature-table borders.
    tbl_pr = signature_table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.append(borders)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    configure_run(p.add_run("MikroFix  |  Omejeno pooblastilo za pregled in prevzem naprav"), size=7.8, color="666666")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
